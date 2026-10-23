#!/usr/bin/env python3
"""
extract_document.py — 任意形式のドキュメントを「位置情報つきプレーンテキスト」に変換する。

レビューでは「どこの話か」を指摘できることが実用性の生命線なので、
本文だけでなく必ず位置マーカー（ページ/スライド/見出し/行/セル番地）を付けて出力する。

使い方:
    python3 extract_document.py <file> [<file> ...] [-o OUTPUT.txt]

対応形式:
    .docx .dotx / .pptx .potx / .xlsx .xlsm .xltx / .csv .tsv
    .pdf / .md .markdown .txt .rst .adoc / .html .htm
    .doc .ppt .xls .odt .odp .ods  (LibreOffice でいったん変換)

出力先を指定しない場合は標準出力。抽出に失敗した要素があっても全体は止めず、
[!抽出失敗] 行を残して続行する（部分的にでも読めたほうが役に立つため）。
"""

from __future__ import annotations

import argparse
import os
import subprocess
import sys
import tempfile

CONVERT_MAP = {
    ".doc": ("docx", ".docx"),
    ".odt": ("docx", ".docx"),
    ".rtf": ("docx", ".docx"),
    ".ppt": ("pptx", ".pptx"),
    ".odp": ("pptx", ".pptx"),
    ".xls": ("xlsx", ".xlsx"),
    ".ods": ("xlsx", ".xlsx"),
}


def warn(msg: str) -> None:
    print(f"[!] {msg}", file=sys.stderr)


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------
def extract_docx(path: str) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(path)
    out: list[str] = []
    para_no = 0
    table_no = 0

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        try:
            if tag == "p":
                para = Paragraph(child, doc)
                para_no += 1
                text = para.text.strip()
                if not text:
                    continue
                style = (para.style.name or "") if para.style is not None else ""
                if style.lower().startswith("heading") or style.startswith("見出し"):
                    digits = "".join(c for c in style if c.isdigit())
                    level = int(digits) if digits else 1
                    out.append(f"\n[¶{para_no}] {'#' * min(level + 1, 6)} {text}")
                elif style in ("Title", "表題"):
                    out.append(f"\n[¶{para_no}] # {text}")
                elif style in ("List Paragraph", "リスト段落"):
                    out.append(f"[¶{para_no}] - {text}")
                else:
                    out.append(f"[¶{para_no}] {text}")
            elif tag == "tbl":
                table_no += 1
                table = Table(child, doc)
                out.append(f"\n[表{table_no}]")
                for r, row in enumerate(table.rows, start=1):
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    out.append(f"[表{table_no}-R{r}] | " + " | ".join(cells) + " |")
                out.append("")
        except Exception as e:  # noqa: BLE001
            out.append(f"[!抽出失敗] {tag}: {e}")

    # 図表のキャプション漏れを防ぐため、画像の存在だけは記録しておく
    n_images = len([r for r in doc.part.rels.values() if "image" in r.reltype])
    if n_images:
        out.append(f"\n[メモ] 埋め込み画像 {n_images} 点あり（本文テキストには含まれない）")
    return "\n".join(out)


# --------------------------------------------------------------------------
# PowerPoint
# --------------------------------------------------------------------------
def extract_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    out: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        layout = ""
        try:
            layout = slide.slide_layout.name
        except Exception:  # noqa: BLE001
            pass
        out.append(f"\n=== [スライド{i}] (レイアウト: {layout}) ===")
        for shape in slide.shapes:
            try:
                if shape.has_table:
                    tbl = shape.table
                    for r, row in enumerate(tbl.rows, start=1):
                        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                        out.append(f"[S{i}-表R{r}] | " + " | ".join(cells) + " |")
                    continue
                if getattr(shape, "has_chart", False) and shape.has_chart:
                    out.append(f"[S{i}-グラフ] {shape.chart.chart_type}")
                    continue
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        name = shape.name or "テキスト"
                        for line in text.split("\n"):
                            if line.strip():
                                out.append(f"[S{i}] ({name}) {line.strip()}")
                elif shape.shape_type is not None:
                    out.append(f"[S{i}] (図形: {shape.name})")
            except Exception as e:  # noqa: BLE001
                out.append(f"[!抽出失敗] S{i}: {e}")
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    out.append(f"[S{i}-ノート] {notes}")
        except Exception:  # noqa: BLE001
            pass
    return "\n".join(out)


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def extract_xlsx(path: str, max_rows: int = 400, max_cols: int = 40) -> str:
    import openpyxl

    out: list[str] = []
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        for ws in wb.worksheets:
            state = getattr(ws, "sheet_state", "visible")
            out.append(f"\n=== [シート: {ws.title}] ({ws.max_row}行 x {ws.max_column}列, 状態={state}) ===")
            rows = 0
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows),
                                    max_col=min(ws.max_column, max_cols)):
                vals = []
                for cell in row:
                    if cell.value is None:
                        continue
                    vals.append(f"{cell.coordinate}={cell.value}")
                if vals:
                    rows += 1
                    out.append(f"[{ws.title}] " + " | ".join(str(v) for v in vals))
            if ws.max_row > max_rows:
                out.append(f"[メモ] {max_rows}行目以降は省略（全{ws.max_row}行）")
            if rows == 0:
                out.append("[メモ] 値のあるセルなし")
    finally:
        wb.close()

    # 数式そのものもレビュー対象になりうるので別途拾う
    try:
        wbf = openpyxl.load_workbook(path, data_only=False)
        formulas = []
        for ws in wbf.worksheets:
            for row in ws.iter_rows(max_row=min(ws.max_row, max_rows),
                                    max_col=min(ws.max_column, max_cols)):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas.append(f"[{ws.title}!{cell.coordinate}] {cell.value}")
        if formulas:
            out.append("\n=== [数式一覧] ===")
            out.extend(formulas[:300])
            if len(formulas) > 300:
                out.append(f"[メモ] 他 {len(formulas) - 300} 件の数式あり")
        wbf.close()
    except Exception as e:  # noqa: BLE001
        warn(f"数式の抽出に失敗: {e}")
    return "\n".join(out)


def extract_csv(path: str, max_rows: int = 400) -> str:
    import csv

    enc = detect_encoding(path)
    delim = "\t" if path.lower().endswith((".tsv", ".tab")) else ","
    out = [f"=== [{os.path.basename(path)}] (encoding={enc}) ==="]
    with open(path, encoding=enc, errors="replace", newline="") as f:
        for i, row in enumerate(csv.reader(f, delimiter=delim), start=1):
            if i > max_rows:
                out.append(f"[メモ] {max_rows}行目以降は省略")
                break
            out.append(f"[R{i}] | " + " | ".join(row) + " |")
    return "\n".join(out)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def extract_pdf(path: str) -> str:
    out: list[str] = []
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                out.append(f"\n=== [p.{i}] ===")
                text = page.extract_text() or ""
                for line in text.split("\n"):
                    if line.strip():
                        out.append(f"[p.{i}] {line.rstrip()}")
                try:
                    for t, table in enumerate(page.extract_tables() or [], start=1):
                        out.append(f"[p.{i}-表{t}]")
                        for r, row in enumerate(table, start=1):
                            cells = ["" if c is None else str(c).replace("\n", " ") for c in row]
                            out.append(f"[p.{i}-表{t}-R{r}] | " + " | ".join(cells) + " |")
                except Exception:  # noqa: BLE001
                    pass
    except Exception as e:  # noqa: BLE001
        warn(f"pdfplumber 失敗({e})。pdftotext にフォールバックします")
        try:
            res = subprocess.run(["pdftotext", "-layout", path, "-"],
                                 capture_output=True, text=True, check=True)
            return res.stdout
        except Exception as e2:  # noqa: BLE001
            return f"[!抽出失敗] PDF を読めませんでした: {e2}"

    body = "\n".join(out)
    if len(body.strip()) < 50:
        body += ("\n[メモ] テキストがほとんど抽出できませんでした。"
                 "スキャン画像 PDF の可能性が高いため、ページを画像化して目視で読むこと。")
    return body


# --------------------------------------------------------------------------
# プレーンテキスト系
# --------------------------------------------------------------------------
def detect_encoding(path: str) -> str:
    try:
        import chardet

        with open(path, "rb") as f:
            raw = f.read(200_000)
        enc = (chardet.detect(raw).get("encoding") or "utf-8")
        return enc
    except Exception:  # noqa: BLE001
        return "utf-8"


def extract_text(path: str) -> str:
    enc = detect_encoding(path)
    with open(path, encoding=enc, errors="replace") as f:
        lines = f.read().split("\n")
    return "\n".join(f"[L{i}] {line.rstrip()}" for i, line in enumerate(lines, start=1))


def extract_html(path: str) -> str:
    from html.parser import HTMLParser

    class Stripper(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self.parts: list[str] = []
            self.skip = False

        def handle_starttag(self, tag, attrs):  # noqa: ANN001
            if tag in ("script", "style"):
                self.skip = True
            if tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
                self.parts.append("\n")
            if tag.startswith("h") and len(tag) == 2 and tag[1].isdigit():
                self.parts.append("#" * int(tag[1]) + " ")

        def handle_endtag(self, tag):  # noqa: ANN001
            if tag in ("script", "style"):
                self.skip = False

        def handle_data(self, data):  # noqa: ANN001
            if not self.skip:
                self.parts.append(data)

    enc = detect_encoding(path)
    with open(path, encoding=enc, errors="replace") as f:
        parser = Stripper()
        parser.feed(f.read())
    text = "".join(parser.parts)
    lines = [ln.strip() for ln in text.split("\n")]
    lines = [ln for ln in lines if ln]
    return "\n".join(f"[L{i}] {ln}" for i, ln in enumerate(lines, start=1))


# --------------------------------------------------------------------------
# ディスパッチ
# --------------------------------------------------------------------------
def convert_with_soffice(path: str) -> str | None:
    ext = os.path.splitext(path)[1].lower()
    if ext not in CONVERT_MAP:
        return None
    target, new_ext = CONVERT_MAP[ext]
    tmpdir = tempfile.mkdtemp(prefix="docrev_")
    try:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", target, "--outdir", tmpdir, path],
            capture_output=True, timeout=180, check=False,
        )
        cand = os.path.join(tmpdir, os.path.splitext(os.path.basename(path))[0] + new_ext)
        return cand if os.path.exists(cand) else None
    except Exception as e:  # noqa: BLE001
        warn(f"soffice 変換に失敗: {e}")
        return None


def extract(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()

    if ext in CONVERT_MAP:
        converted = convert_with_soffice(path)
        if converted:
            warn(f"{ext} を {os.path.splitext(converted)[1]} に変換して処理します")
            path, ext = converted, os.path.splitext(converted)[1].lower()
        else:
            return f"[!抽出失敗] {ext} の変換に失敗しました。手動で確認してください。"

    if ext in (".docx", ".dotx", ".docm"):
        return extract_docx(path)
    if ext in (".pptx", ".potx", ".pptm"):
        return extract_pptx(path)
    if ext in (".xlsx", ".xlsm", ".xltx"):
        return extract_xlsx(path)
    if ext in (".csv", ".tsv", ".tab"):
        return extract_csv(path)
    if ext == ".pdf":
        return extract_pdf(path)
    if ext in (".html", ".htm", ".xhtml"):
        return extract_html(path)
    return extract_text(path)


def main() -> int:
    ap = argparse.ArgumentParser(description="ドキュメントを位置情報つきテキストに変換する")
    ap.add_argument("files", nargs="+")
    ap.add_argument("-o", "--output", help="出力先ファイル（省略時は標準出力）")
    args = ap.parse_args()

    chunks: list[str] = []
    for path in args.files:
        if not os.path.exists(path):
            chunks.append(f"[!] ファイルが見つかりません: {path}")
            continue
        size = os.path.getsize(path)
        header = (f"\n{'=' * 70}\n"
                  f"FILE: {os.path.basename(path)}  ({size:,} bytes)\n"
                  f"PATH: {path}\n"
                  f"{'=' * 70}")
        try:
            body = extract(path)
        except Exception as e:  # noqa: BLE001
            body = f"[!抽出失敗] {type(e).__name__}: {e}"
        chars = len(body)
        chunks.append(f"{header}\n[統計] 抽出テキスト {chars:,} 文字\n{body}")

    result = "\n".join(chunks)
    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(result)
        print(f"書き出しました: {args.output} ({len(result):,} 文字)")
    else:
        print(result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
